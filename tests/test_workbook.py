"""The generated workbook itself: README text vs the Checks sheet, the exchequer bottom line,
and (when a formula evaluator is available) that every consistency check evaluates to PASS."""
import os
import runpy

import openpyxl
import pytest

SCRIPT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts", "build_model.py")


@pytest.fixture(scope="module")
def xlsx(tmp_path_factory):
    work = tmp_path_factory.mktemp("wb")
    (work / "outputs").mkdir()
    prev = os.getcwd()
    os.chdir(work)
    try:
        ns = runpy.run_path(SCRIPT, run_name="__main__")
    finally:
        os.chdir(prev)
    return str(work / ns["OUT"]), ns


def all_text(sheets):
    return [c.value for ws in sheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def test_readme_check_count_matches_checks_sheet(xlsx):
    path, ns = xlsx
    wb = openpyxl.load_workbook(path)
    n = len(ns["checks"])
    line = next(t for t in all_text(wb) if t.startswith("7. Checks:"))
    assert f"{n} automated consistency checks" in line


def test_exchequer_bottom_line_does_not_hardcode_zero(xlsx):
    path, _ = xlsx
    wb = openpyxl.load_workbook(path)
    line = next(t for t in all_text([wb["ExchequerRevenue"]]) if "Bottom line" in t)
    assert "~Rs 0 " not in line


@pytest.fixture(scope="module")
def evaluator(xlsx):
    pycel = pytest.importorskip("pycel")
    xl = pycel.ExcelCompiler(filename=xlsx[0])
    try:
        xl.evaluate("Checks!B5")
    except Exception as e:  # pycel 1.0b30 cannot compile formulas on Python 3.14 (ast.Num removed)
        pytest.skip(f"pycel cannot evaluate here: {type(e).__name__}")
    return xl


def test_every_consistency_check_passes(xlsx, evaluator):
    path, ns = xlsx
    ws = openpyxl.load_workbook(path)["Checks"]
    labels = {ws.cell(r, 1).value: r for r in range(5, 5 + len(ns["checks"]))}
    results = {lab: evaluator.evaluate(f"Checks!B{r}") for lab, r in labels.items()}
    assert results == {lab: "PASS" for lab in labels}


def test_recurring_fuel_tax_hand_computed(xlsx, evaluator):
    """(42.663 MMT/0.74 kg/L x Rs 31 + 93.946 MMT/0.83 x Rs 22) in lakh crore."""
    _, ns = xlsx
    petrol_cr_l = 42.663 * 100 / 0.74
    diesel_cr_l = 93.946 * 100 / 0.83
    expected = (petrol_cr_l * (13 + 18) + diesel_cr_l * (10 + 12)) / 1e5
    assert evaluator.evaluate(ns["EXC"]["fuel_total"].replace("$", "")) == pytest.approx(expected, rel=1e-9)
